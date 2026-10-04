from docx import Document

document = Document()

document.add_heading("KnowledgeBot Document Test", level=1)

document.add_paragraph(
    "This is a test DOCX document."
)

document.add_paragraph(
    "We are testing DOCX text extraction."
)

document.add_paragraph(
    "The document service should extract this text."
)

document.save("test.docx")

print("test.docx created successfully")